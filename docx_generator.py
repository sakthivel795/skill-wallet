from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH
)
from docx.enum.table import (
    WD_TABLE_ALIGNMENT
)
from docx.shared import (
    Inches,
    Pt
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from document_utils.text_utils import (
    sanitize_text
)


BASE_DIR = Path(__file__).resolve().parent.parent
LOGO_PATH = BASE_DIR / "assets" / "logo.png"


def add_page_number(paragraph):

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()

    field_begin = OxmlElement(
        "w:fldChar"
    )

    field_begin.set(
        qn("w:fldCharType"),
        "begin"
    )

    instruction = OxmlElement(
        "w:instrText"
    )

    instruction.set(
        qn("xml:space"),
        "preserve"
    )

    instruction.text = "PAGE"

    field_end = OxmlElement(
        "w:fldChar"
    )

    field_end.set(
        qn("w:fldCharType"),
        "end"
    )

    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def format_docx(
    text: str,
    doc_type: str,
    terms=None
) -> bytes:

    text = sanitize_text(text)

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Default font
    styles = document.styles

    normal_style = styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    # Logo
    if LOGO_PATH.exists():

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run()

        run.add_picture(
            str(LOGO_PATH),
            width=Inches(1.5)
        )

    # Title
    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        sanitize_text(doc_type).upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    # Main content
    lines = text.splitlines()

    for line in lines:

        clean_line = sanitize_text(line)

        if not clean_line:
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15

        # Section headings
        if (
            clean_line.isupper()
            or clean_line.endswith(":")
        ):

            run = paragraph.add_run(
                clean_line
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        elif clean_line.startswith(
            ("1.", "2.", "3.", "4.", "5.",
             "6.", "7.", "8.", "9.")
        ):

            run = paragraph.add_run(
                clean_line
            )

            run.bold = True
            run.font.name = "Times New Roman"

        elif clean_line.startswith(
            ("-", "*", "•")
        ):

            run = paragraph.add_run(
                clean_line
            )

            run.font.name = "Times New Roman"

        else:

            run = paragraph.add_run(
                clean_line
            )

            run.font.name = "Times New Roman"

    # Terms table
    if terms:

        document.add_paragraph()

        heading = document.add_paragraph()

        run = heading.add_run(
            "KEY TERMS"
        )

        run.bold = True
        run.font.size = Pt(12)

        table = document.add_table(
            rows=1,
            cols=2
        )

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        table.style = "Table Grid"

        header = table.rows[0].cells

        header[0].text = "No."
        header[1].text = "Term"

        for index, term in enumerate(
            terms,
            start=1
        ):

            row = table.add_row().cells

            row[0].text = str(index)
            row[1].text = sanitize_text(term)

    # Footer
    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "LegalEase - AI-assisted legal document draft | Page "
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    add_page_number(
        footer_paragraph
    )

    # Save to memory
    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()